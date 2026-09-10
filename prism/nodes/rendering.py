import os
import docx
from docx.shared import Pt
from prism.state import PRISMState
from prism.config import OUTPUTS_DIR
from prism.nodes.specialists import pptx_renderer_node

def _create_docx(title: str, text: str, output_path: str):
    doc = docx.Document()
    
    # Title
    heading = doc.add_heading(title, level=0)
    
    # Content paragraphs
    lines = text.split("\n")
    for line in lines:
        line_s = line.strip()
        if not line_s:
            continue
        if line_s.startswith("### "):
            doc.add_heading(line_s.replace("### ", ""), level=3)
        elif line_s.startswith("## "):
            doc.add_heading(line_s.replace("## ", ""), level=2)
        elif line_s.startswith("# "):
            doc.add_heading(line_s.replace("# ", ""), level=1)
        elif line_s.startswith("- ") or line_s.startswith("* "):
            doc.add_paragraph(line_s[2:], style='List Bullet')
        else:
            p = doc.add_paragraph(line_s)
            p.style.font.size = Pt(11)
            
    doc.save(output_path)

def render_node(state: PRISMState) -> dict:
    run_id = state.get("run_id", "latest")
    selected_outputs = state.get("selected_outputs", [])

    rendered_assets = dict(state.get("rendered_assets") or {})
    pptx_path = state.get("pptx_path")

    # Render presentation if selected
    if "presentation" in selected_outputs and state.get("presentation_output"):
        ppt_updates = pptx_renderer_node(state)
        pptx_path = ppt_updates.get("pptx_path", pptx_path)
        if ppt_updates.get("rendered_assets"):
            rendered_assets.update(ppt_updates["rendered_assets"])

    # Render LinkedIn
    if "linkedin" in selected_outputs and state.get("linkedin_output"):
        content = state["linkedin_output"]
        md_path = os.path.join(OUTPUTS_DIR, f"linkedin_{run_id}.md")
        txt_path = os.path.join(OUTPUTS_DIR, f"linkedin_{run_id}.txt")
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(f"# PRISM Generated LinkedIn Post\n\n{content}\n")
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write(content)
        rendered_assets["linkedin_md"] = md_path
        rendered_assets["linkedin_txt"] = txt_path

    # Render Twitter/X
    if "twitter" in selected_outputs and state.get("twitter_output"):
        content = state["twitter_output"]
        md_path = os.path.join(OUTPUTS_DIR, f"twitter_{run_id}.md")
        txt_path = os.path.join(OUTPUTS_DIR, f"twitter_{run_id}.txt")
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(f"# PRISM Generated Twitter/X Post\n\n{content}\n")
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write(content)
        rendered_assets["twitter_md"] = md_path
        rendered_assets["twitter_txt"] = txt_path

    # Render Executive Summary
    if "summary" in selected_outputs and state.get("summary_output"):
        content = state["summary_output"]
        docx_path = os.path.join(OUTPUTS_DIR, f"executive_summary_{run_id}.docx")
        md_path = os.path.join(OUTPUTS_DIR, f"executive_summary_{run_id}.md")
        _create_docx("PRISM Executive Summary", content, docx_path)
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(f"# Executive Summary\n\n{content}\n")
        rendered_assets["summary_docx"] = docx_path
        rendered_assets["summary_md"] = md_path

    # Render Advisory Report
    if "advisory" in selected_outputs and state.get("advisory_output"):
        content = state["advisory_output"]
        docx_path = os.path.join(OUTPUTS_DIR, f"advisory_report_{run_id}.docx")
        md_path = os.path.join(OUTPUTS_DIR, f"advisory_report_{run_id}.md")
        _create_docx("PRISM Strategic Advisory Report", content, docx_path)
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(f"# Strategic Advisory Report\n\n{content}\n")
        rendered_assets["advisory_docx"] = docx_path
        rendered_assets["advisory_md"] = md_path

    return {"rendered_assets": rendered_assets, "pptx_path": pptx_path}
