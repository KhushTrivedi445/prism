import os
from pptx import Presentation as PPTXPresentation
from pptx.util import Pt
from prism.state import PRISMState
from prism.schemas import Presentation
from prism.config import get_llm, OUTPUTS_DIR
from prism.utils import invoke_structured_llm, invoke_text_llm

def linkedin_node(state: PRISMState) -> dict:
    fact_graph = state.get("fact_graph")
    tone = state.get("tone", "professional")

    prompt = f"""
    You are the LinkedIn Content Specialist Agent for PRISM.

    Your task is to create a professional LinkedIn post
    using ONLY the facts provided in the Fact Graph.

    Rules:
    1. Use only information supported by the Fact Graph.
    2. Do not invent statistics, features, achievements, or claims.
    3. Do not change the meaning of any fact.
    4. Make the post engaging and suitable for LinkedIn.
    5. Use the requested tone.
    6. You may use emojis and hashtags when appropriate.
    7. Do not mention the Fact Graph in the final post.
    8. Return only the final LinkedIn post.

    Tone:
    {tone}

    Fact Graph:
    {fact_graph}
    """

    content = invoke_text_llm(prompt)
    return {"linkedin_output": content}


def twitter_node(state: PRISMState) -> dict:
    fact_graph = state.get("fact_graph")
    tone = state.get("tone", "engaging")

    prompt = f"""
    You are the Twitter/X Content Specialist Agent for PRISM.

    Your task is to create a concise and engaging Twitter/X post
    using ONLY the facts provided in the Fact Graph.

    Rules:
    1. Use only information supported by the Fact Graph.
    2. Do not invent statistics, features, achievements, or claims.
    3. Do not change the meaning of any fact.
    4. Keep the post concise and suitable for Twitter/X.
    5. Start with an engaging hook.
    6. Use the requested tone.
    7. Use a small number of relevant hashtags.
    8. Do not mention the Fact Graph in the final post.
    9. Return only the final Twitter/X post.

    Tone:
    {tone}

    Fact Graph:
    {fact_graph}
    """

    content = invoke_text_llm(prompt)
    return {"twitter_output": content}


def summary_node(state: PRISMState) -> dict:
    fact_graph = state.get("fact_graph")
    tone = state.get("tone", "executive")

    prompt = f"""
    You are the Executive Summary Specialist Agent for PRISM.

    Your task is to create a concise but sufficiently detailed
    executive summary using ONLY the facts provided in the Fact Graph.

    Rules:
    1. Use only information supported by the Fact Graph.
    2. Do not invent statistics, features, achievements, or claims.
    3. Do not change the meaning of any fact.
    4. Cover the key purpose, functionality, and important relationships
       present in the Fact Graph.
    5. Give enough context for a reader to understand the system without
       reading the original source.
    6. Aim for approximately 80–150 words when enough information is available.
    7. Do not add length by repeating the same information.
    8. Use a professional and clear writing style.
    9. Follow the requested tone.
    10. Do not mention the Fact Graph in the final summary.
    11. Return only the final executive summary.

    Tone:
    {tone}

    Fact Graph:
    {fact_graph}
    """

    content = invoke_text_llm(prompt)
    return {"summary_output": content}


def advisory_node(state: PRISMState) -> dict:
    fact_graph = state.get("fact_graph")
    tone = state.get("tone", "formal")

    prompt = f"""
    You are the Advisory Specialist Agent for PRISM.

    Your task is to provide clear and useful advisory guidance
    based ONLY on the information contained in the Fact Graph.

    Rules:
    1. Use only information supported by the Fact Graph.
    2. Do not invent facts, statistics, features, or capabilities.
    3. Do not make claims that cannot be supported by the Fact Graph.
    4. Identify practical implications or recommendations that
       logically follow from the provided facts.
    5. Clearly distinguish recommendations from factual statements.
    6. Keep the advice useful, specific, and easy to understand.
    7. Do not repeat the entire Fact Graph.
    8. Follow the requested tone.
    9. Do not mention the Fact Graph in the final response.
    10. Return only the final advisory content.

    Tone:
    {tone}

    Fact Graph:
    {fact_graph}
    """

    content = invoke_text_llm(prompt)
    return {"advisory_output": content}


def presentation_node(state: PRISMState) -> dict:
    fact_graph = state.get("fact_graph")
    tone = state.get("tone", "professional")
    llm = get_llm()

    prompt = f"""
    You are the Presentation Specialist Agent for PRISM.

    Your task is to create a clear presentation structure
    using ONLY the information contained in the Fact Graph.

    Create a logical presentation that explains the important
    information in a way that is easy to understand.

    Rules:
    1. Use only facts supported by the Fact Graph.
    2. Do not invent statistics, features, achievements, or claims.
    3. Do not change the meaning of any fact.
    4. Create a logical sequence of slides.
    5. Each slide should have a clear title.
    6. Use concise bullet points.
    7. Avoid putting too much text on one slide.
    8. Focus on important information rather than repeating facts.
    9. Follow the requested tone.
    10. Do not include speaker notes.
    11. Return the presentation as valid JSON matching the Presentation schema.

    Suggested structure when enough information is available:
    - Introduction / Title
    - Overview
    - How it works
    - Important components
    - Key value or implications
    - Conclusion

    Do not force unnecessary slides if the Fact Graph
    does not contain enough information.

    Tone:
    {tone}

    Fact Graph:
    {fact_graph}
    """

    presentation = invoke_structured_llm(llm, Presentation, prompt)

    if isinstance(presentation, dict):
        pres_dict = presentation
    else:
        pres_dict = presentation.model_dump()

    return {"presentation_output": pres_dict}


def pptx_renderer_node(state: PRISMState) -> dict:
    presentation_data = state.get("presentation_output")
    if not presentation_data:
        return {}

    prs = PPTXPresentation()
    title = presentation_data.get("title", "PRISM Presentation")
    slides = presentation_data.get("slides", [])

    for index, slide_data in enumerate(slides):
        if index == 0:
            layout = prs.slide_layouts[0]
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = slide_data.get("title", title)

            if len(slide.placeholders) > 1:
                subtitle = slide.placeholders[1]
                subtitle.text = "\n".join(slide_data.get("content", []))
        else:
            layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = slide_data.get("title", f"Slide {index+1}")

            if len(slide.placeholders) > 1:
                text_frame = slide.placeholders[1].text_frame
                text_frame.clear()
                for bullet in slide_data.get("content", []):
                    paragraph = text_frame.add_paragraph()
                    paragraph.text = bullet
                    paragraph.level = 0
                    paragraph.font.size = Pt(22)

    run_id = state.get("run_id", "default")
    output_path = os.path.join(OUTPUTS_DIR, f"PRISM_Presentation_{run_id}.pptx")
    prs.save(output_path)

    rendered_assets = dict(state.get("rendered_assets") or {})
    rendered_assets["presentation"] = output_path

    return {"pptx_path": output_path, "rendered_assets": rendered_assets}
